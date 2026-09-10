"""Generate CommoHedge commercial pitch deck (PowerPoint)."""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import nsmap
from pptx.oxml import parse_xml
from pptx.util import Inches, Pt, Emu

# Brand colors (landing terminal)
NAVY = RGBColor(0x0C, 0x13, 0x22)
NAVY_DEEP = RGBColor(0x07, 0x0E, 0x1D)
NAVY_CARD = RGBColor(0x14, 0x1B, 0x2B)
LIME = RGBColor(0xAE, 0xF8, 0x33)
LIME_DARK = RGBColor(0x21, 0x36, 0x00)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MUTED = RGBColor(0xC1, 0xCA, 0xAF)
SOFT = RGBColor(0xDC, 0xE2, 0xF7)
BORDER = RGBColor(0x42, 0x4A, 0x35)
RED_SOFT = RGBColor(0xF8, 0x71, 0x71)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

OUT = Path(__file__).resolve().parent / "CommoHedge_Presentation_Commerciale.pptx"


def set_run(run, size=18, bold=False, color=WHITE, font="Calibri"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def add_text_box(slide, left, top, width, height, text, size=18, bold=False, color=WHITE, align=PP_ALIGN.LEFT, font="Calibri"):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_run(run, size=size, bold=bold, color=color, font=font)
    return box


def add_para(tf, text, size=14, bold=False, color=SOFT, space_before=6, space_after=2, align=PP_ALIGN.LEFT):
    p = tf.add_paragraph()
    p.alignment = align
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    run = p.add_run()
    run.text = text
    set_run(run, size=size, bold=blank if False else bold, color=color)
    return p


def fill_shape(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def card(slide, left, top, width, height, fill=NAVY_CARD):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    fill_shape(shape, fill)
    shape.adjustments[0] = 0.08
    shape.line.color.rgb = BORDER
    shape.line.width = Pt(1)
    return shape


def accent_bar(slide, left, top, width, height=Inches(0.08)):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    fill_shape(bar, LIME)
    return bar


def footer(slide, page, total=14):
    add_text_box(
        slide,
        Inches(0.5),
        Inches(7.1),
        Inches(8),
        Inches(0.3),
        "COMMOHEDGE  ·  Commodity hedging & intelligence terminal  ·  Confidentiel",
        size=10,
        color=MUTED,
    )
    add_text_box(
        slide,
        Inches(11.5),
        Inches(7.1),
        Inches(1.3),
        Inches(0.3),
        f"{page} / {total}",
        size=10,
        color=MUTED,
        align=PP_ALIGN.RIGHT,
    )


def bg(slide):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    fill_shape(shape, NAVY_DEEP)
    # keep behind
    spTree = slide.shapes._spTree
    sp = shape._element
    spTree.remove(sp)
    spTree.insert(2, sp)


def section_title(slide, title, subtitle=None):
    accent_bar(slide, Inches(0.55), Inches(0.42), Inches(0.9), Inches(0.07))
    add_text_box(slide, Inches(0.55), Inches(0.55), Inches(12), Inches(0.55), title, size=28, bold=True, color=WHITE)
    if subtitle:
        add_text_box(slide, Inches(0.55), Inches(1.1), Inches(12), Inches(0.4), subtitle, size=14, color=MUTED)


def new_slide(prs):
    blank = prs.slide_layouts[6]  # blank
    slide = prs.slides.add_slide(blank)
    bg(slide)
    return slide


def build():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    total = 14

    # ── 1. Cover ──────────────────────────────────────────────
    s = new_slide(prs)
    # left lime strip
    strip = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.18), SLIDE_H)
    fill_shape(strip, LIME)

    add_text_box(s, Inches(0.7), Inches(1.6), Inches(11), Inches(0.4), "PITCH COMMERCIAL  ·  2026", size=14, bold=True, color=LIME)
    add_text_box(s, Inches(0.7), Inches(2.15), Inches(11), Inches(1.0), "COMMOHEDGE", size=54, bold=True, color=WHITE)
    add_text_box(
        s,
        Inches(0.7),
        Inches(3.15),
        Inches(11),
        Inches(0.5),
        "Commodity hedging & intelligence terminal",
        size=22,
        color=LIME,
    )
    add_text_box(
        s,
        Inches(0.7),
        Inches(3.85),
        Inches(10.5),
        Inches(1.0),
        "Pricer, couvrir, monitorer et contextualiser le risque matières premières\n"
        "(et FX) dans une workstation professionnelle unique.",
        size=16,
        color=SOFT,
    )
    add_text_box(s, Inches(0.7), Inches(5.4), Inches(10), Inches(0.35), "Pour trésoriers corporate · desks commodity · brokers · shipping", size=13, color=MUTED)
    footer(s, 1, total)

    # ── 2. Agenda ─────────────────────────────────────────────
    s = new_slide(prs)
    section_title(s, "Agenda", "Structure de la présentation")
    items = [
        ("01", "Le problème"),
        ("02", "La solution CommoHedge"),
        ("03", "Modules & fonctionnalités"),
        ("04", "Instruments & modèles"),
        ("05", "Cibles & différenciation"),
        ("06", "Parcours démo & packaging"),
        ("07", "Closing & prochaines étapes"),
    ]
    for i, (num, label) in enumerate(items):
        y = Inches(1.7) + Inches(i * 0.65)
        nbox = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.7), y, Inches(0.7), Inches(0.48))
        fill_shape(nbox, LIME)
        nbox.adjustments[0] = 0.2
        add_text_box(s, Inches(0.7), y + Inches(0.08), Inches(0.7), Inches(0.35), num, size=14, bold=True, color=LIME_DARK, align=PP_ALIGN.CENTER)
        add_text_box(s, Inches(1.65), y + Inches(0.08), Inches(8), Inches(0.4), label, size=18, bold=True, color=WHITE)
    footer(s, 2, total)

    # ── 3. Problem ────────────────────────────────────────────
    s = new_slide(prs)
    section_title(s, "Le problème", "Pourquoi les équipes mid-market restent fragiles")
    problems = [
        ("Coût des terminaux", "Bloomberg / Murex / stacks quant = budget et IT hors de portée pour beaucoup de corporates."),
        ("Outils fragmentés", "Pricing Excel, expositions ailleurs, news ailleurs, freight ailleurs — aucune boucle unifiée."),
        ("Risque opaque", "Barriers & digitals mal pricés, VaR approximatif, pas de lien stratégie → book → comité."),
    ]
    for i, (title, body) in enumerate(problems):
        x = Inches(0.55) + Inches(i * 4.15)
        card(s, x, Inches(1.85), Inches(3.95), Inches(3.8))
        accent_bar(s, x + Inches(0.3), Inches(2.15), Inches(0.7))
        add_text_box(s, x + Inches(0.3), Inches(2.45), Inches(3.35), Inches(0.6), title, size=18, bold=True, color=WHITE)
        add_text_box(s, x + Inches(0.3), Inches(3.2), Inches(3.35), Inches(2.0), body, size=14, color=SOFT)
    footer(s, 3, total)

    # ── 4. Solution / pitch ───────────────────────────────────
    s = new_slide(prs)
    section_title(s, "La solution", "Une phrase d'accroche pour ouvrir le meeting")
    card(s, Inches(0.55), Inches(1.85), Inches(12.2), Inches(2.4))
    add_text_box(
        s,
        Inches(0.9),
        Inches(2.15),
        Inches(11.5),
        Inches(1.9),
        "CommoHedge remplace le patchwork Excel + Bloomberg + mails banque\n"
        "par un terminal unique : pricing desk-grade, book de couverture avec MTM,\n"
        "data marché (futures / vol / freight) et intelligence géopolitique.",
        size=20,
        bold=True,
        color=WHITE,
    )

    pillars = [("PRICE", "Black-76, barriers, MC, Greeks"), ("HEDGE", "Exposures → book → MTM"), ("MEASURE", "VaR, stress, reports PDF"), ("SENSE", "Data + Intel + AI")]
    for i, (t, b) in enumerate(pillars):
        x = Inches(0.55) + Inches(i * 3.15)
        c = card(s, x, Inches(4.55), Inches(3.0), Inches(1.7))
        add_text_box(s, x + Inches(0.2), Inches(4.75), Inches(2.6), Inches(0.4), t, size=16, bold=True, color=LIME)
        add_text_box(s, x + Inches(0.2), Inches(5.25), Inches(2.6), Inches(0.7), b, size=13, color=SOFT)
    footer(s, 4, total)

    # ── 5. Value props ────────────────────────────────────────
    s = new_slide(prs)
    section_title(s, "Proposition de valeur", "Ce que gagne le client")
    vals = [
        ("Réduire", "Coût & dépendance", "Moins de licences terminal et de spreadsheets non gouvernés."),
        ("Accélérer", "Décision de couverture", "De l'idée de structure au PDF comité en une session."),
        ("Sécuriser", "Gouvernance risque", "Rôles, sync cloud, MTM, VaR / stress traçables."),
    ]
    for i, (k, t, b) in enumerate(vals):
        y = Inches(1.8) + Inches(i * 1.55)
        card(s, Inches(0.55), y, Inches(12.2), Inches(1.4))
        badge = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.85), y + Inches(0.35), Inches(1.8), Inches(0.55))
        fill_shape(badge, LIME)
        badge.adjustments[0] = 0.2
        add_text_box(s, Inches(0.85), y + Inches(0.42), Inches(1.8), Inches(0.4), k, size=14, bold=True, color=LIME_DARK, align=PP_ALIGN.CENTER)
        add_text_box(s, Inches(3.0), y + Inches(0.3), Inches(9), Inches(0.4), t, size=18, bold=True, color=WHITE)
        add_text_box(s, Inches(3.0), y + Inches(0.75), Inches(9), Inches(0.45), b, size=14, color=SOFT)
    footer(s, 5, total)

    # ── 6. Modules map ────────────────────────────────────────
    s = new_slide(prs)
    section_title(s, "Cartographie produit", "6 modules — 20+ pages métier")
    modules = [
        ("Pricing & Strategy", "Pricers, Strategy Builder", "Structures multi-jambes, stress, backtest"),
        ("Risk Operations", "Exposures, Hedging, Positions", "Boucle exposition → couverture → MTM"),
        ("Risk Analytics", "Risk, Regression, Reports", "VaR / ES, stress, export PDF & CSV"),
        ("Market Data", "Market, Terminal, Rates", "Futures, options, vol surface 3D, courbes"),
        ("Intelligence", "Intel Workspace, Map, News", "Géopolitique, AIS, macro, freight context"),
        ("AI & Admin", "Hedge Assistant, Users, Settings", "Conseil IA, rôles, white-label, sync"),
    ]
    for i, (title, pages, value) in enumerate(modules):
        col = i % 3
        row = i // 3
        x = Inches(0.55) + Inches(col * 4.15)
        y = Inches(1.75) + Inches(row * 2.4)
        card(s, x, y, Inches(3.95), Inches(2.2))
        add_text_box(s, x + Inches(0.25), y + Inches(0.25), Inches(3.4), Inches(0.4), title, size=15, bold=True, color=LIME)
        add_text_box(s, x + Inches(0.25), y + Inches(0.75), Inches(3.4), Inches(0.4), pages, size=12, bold=True, color=WHITE)
        add_text_box(s, x + Inches(0.25), y + Inches(1.2), Inches(3.4), Inches(0.7), value, size=12, color=SOFT)
    footer(s, 6, total)

    # ── 7. Instruments ────────────────────────────────────────
    s = new_slide(prs)
    section_title(s, "Instruments & modèles", "Crédibilité desk — au-delà d'Excel")
    left_items = [
        "Forwards & Swaps",
        "Vanillas Call / Put",
        "Barriers KO / KI (single, reverse, double)",
        "Digitals & touches (one-touch, no-touch, range)",
        "Structures multi-jambes & zero-cost",
    ]
    right_items = [
        "Black-76 (options sur forward)",
        "Black-Scholes & Garman-Kohlhagen (FX)",
        "Barriers closed-form + Monte Carlo",
        "Greeks : Δ Γ Θ Vega Rho + IV",
        "Cost of carry · courbes de taux",
    ]
    card(s, Inches(0.55), Inches(1.75), Inches(6.0), Inches(4.5))
    add_text_box(s, Inches(0.85), Inches(2.0), Inches(5.4), Inches(0.4), "Catalogue instruments", size=16, bold=True, color=LIME)
    for i, t in enumerate(left_items):
        add_text_box(s, Inches(0.85), Inches(2.6) + Inches(i * 0.55), Inches(5.4), Inches(0.45), f"▸  {t}", size=14, color=SOFT)

    card(s, Inches(6.8), Inches(1.75), Inches(6.0), Inches(4.5))
    add_text_box(s, Inches(7.1), Inches(2.0), Inches(5.4), Inches(0.4), "Moteur de pricing", size=16, bold=True, color=LIME)
    for i, t in enumerate(right_items):
        add_text_box(s, Inches(7.1), Inches(2.6) + Inches(i * 0.55), Inches(5.4), Inches(0.45), f"▸  {t}", size=14, color=SOFT)
    footer(s, 7, total)

    # ── 8. Underlyings ────────────────────────────────────────
    s = new_slide(prs)
    section_title(s, "Univers couvert", "26+ commodités · FX · rates · freight")
    cats = [
        ("ENERGY", "WTI · Brent · NatGas · Heating Oil · RBOB"),
        ("METALS", "Gold · Silver · Platinum · Palladium · Cu · Al · Zn · Ni"),
        ("AGRI / LIVESTOCK", "Corn · Wheat · Soy · Coffee · Sugar · Cotton · Cattle · Hogs"),
        ("FREIGHT / BUNKER", "Baltic / routes container & dirty · Ship & Bunker"),
        ("FX & RATES", "Garman-Kohlhagen · USD EUR GBP JPY CHF CAD SGD"),
    ]
    for i, (title, body) in enumerate(cats):
        y = Inches(1.7) + Inches(i * 0.95)
        card(s, Inches(0.55), y, Inches(12.2), Inches(0.85))
        add_text_box(s, Inches(0.85), y + Inches(0.22), Inches(2.8), Inches(0.4), title, size=13, bold=True, color=LIME)
        add_text_box(s, Inches(3.8), y + Inches(0.22), Inches(8.6), Inches(0.45), body, size=14, color=SOFT)
    footer(s, 8, total)

    # ── 9. Target buyers ──────────────────────────────────────
    s = new_slide(prs)
    section_title(s, "À qui vendre", "Segments prioritaires & angles de pitch")
    buyers = [
        ("Trésorerie corporate", "Couvrir énergie / métaux / agri sans desk banque complet", "Workstation Bloomberg-light, décision en minutes"),
        ("Trading houses / brokers", "Pricers client-facing + book de couverture", "Terminal desk + packs PDF commerciaux"),
        ("Freight & shipping", "Bunker, freight et énergie sur le même écran", "Différenciateur rare vs outils metals/oil seuls"),
        ("Cabinets / white-label", "Plateforme marque client + gouvernance", "Logo, company settings, rôles Admin → Viewer"),
    ]
    # header row
    add_text_box(s, Inches(0.7), Inches(1.65), Inches(3.2), Inches(0.35), "SEGMENT", size=11, bold=True, color=MUTED)
    add_text_box(s, Inches(4.0), Inches(1.65), Inches(4.5), Inches(0.35), "BESOIN", size=11, bold=True, color=MUTED)
    add_text_box(s, Inches(8.6), Inches(1.65), Inches(4.2), Inches(0.35), "ANGLE PITCH", size=11, bold=True, color=MUTED)
    for i, (seg, need, pitch) in enumerate(buyers):
        y = Inches(2.05) + Inches(i * 1.1)
        card(s, Inches(0.55), y, Inches(12.2), Inches(1.0))
        add_text_box(s, Inches(0.75), y + Inches(0.28), Inches(3.1), Inches(0.5), seg, size=14, bold=True, color=WHITE)
        add_text_box(s, Inches(4.0), y + Inches(0.28), Inches(4.4), Inches(0.55), need, size=13, color=SOFT)
        add_text_box(s, Inches(8.6), y + Inches(0.28), Inches(3.9), Inches(0.55), pitch, size=13, color=LIME)
    footer(s, 9, total)

    # ── 10. Competition ───────────────────────────────────────
    s = new_slide(prs)
    section_title(s, "Différenciation", "Pourquoi CommoHedge gagne le mid-market")
    comps = [
        ("vs Bloomberg / Refinitiv", "Coût & complexité trop élevés", "80% de la valeur hedge desk à une fraction du coût"),
        ("vs Excel + VBA", "Fragile, non auditable, pas de book", "Modèles institutionnels + workflow expositions/MTM"),
        ("vs Pricers bancaires", "Opaque, lent, siloté FX vs commodités", "ADN FX + commodités dans un seul terminal"),
        ("vs Outils risk purs", "Analytique sans data ni intel", "Pricing + data + intel + AI intégrés"),
    ]
    for i, (vs, limit, edge) in enumerate(comps):
        y = Inches(1.7) + Inches(i * 1.15)
        card(s, Inches(0.55), y, Inches(12.2), Inches(1.05))
        add_text_box(s, Inches(0.8), y + Inches(0.3), Inches(3.5), Inches(0.45), vs, size=14, bold=True, color=WHITE)
        add_text_box(s, Inches(4.4), y + Inches(0.3), Inches(3.8), Inches(0.5), limit, size=13, color=MUTED)
        add_text_box(s, Inches(8.4), y + Inches(0.3), Inches(4.1), Inches(0.5), edge, size=13, color=LIME)
    footer(s, 10, total)

    # ── 11. Demo script ───────────────────────────────────────
    s = new_slide(prs)
    section_title(s, "Script démo (20–25 min)", "Ordre recommandé pour convaincre")
    steps = [
        ("01", "Landing + Request Access", "Positionnement terminal + funnel"),
        ("02", "Commodity Market / Data Terminal", "Crédibilité data : spots, futures, vol"),
        ("03", "Strategy Builder + Pricers", "Cœur produit : structure & payoff"),
        ("04", "Exposures → Hedging → MTM", "Preuve opérationnelle"),
        ("05", "Risk Analysis + Reports PDF", "Langage comité / direction"),
        ("06", "Intel Workspace + Hedge Assistant", "Wow factor géopolitique + IA"),
    ]
    for i, (num, screen, goal) in enumerate(steps):
        col = i % 3
        row = i // 3
        x = Inches(0.55) + Inches(col * 4.15)
        y = Inches(1.8) + Inches(row * 2.35)
        card(s, x, y, Inches(3.95), Inches(2.15))
        add_text_box(s, x + Inches(0.25), y + Inches(0.25), Inches(3.4), Inches(0.4), num, size=22, bold=True, color=LIME)
        add_text_box(s, x + Inches(0.25), y + Inches(0.8), Inches(3.4), Inches(0.55), screen, size=14, bold=True, color=WHITE)
        add_text_box(s, x + Inches(0.25), y + Inches(1.4), Inches(3.4), Inches(0.5), goal, size=12, color=SOFT)
    footer(s, 11, total)

    # ── 12. Packaging ─────────────────────────────────────────
    s = new_slide(prs)
    section_title(s, "Packaging commercial", "3 offres pour cadrer la négociation")
    packs = [
        ("ESSENTIEL", "Core Hedge", "Dashboard, Exposures, Hedging,\nStrategy Builder, Pricers, Risk,\nReports.\n\nIdéal trésorerie.", False),
        ("RECOMMANDÉ", "Desk Pro", "Core +\nCommodity Market +\nData Terminal +\nRate Explorer +\nHedge Assistant.\n\nPour desks & brokers.", True),
        ("PREMIUM", "Intelligence", "Desk Pro +\nIntel Workspace,\nWorld Map, news,\nwhite-label,\nSSO / admin avancé.", False),
    ]
    for i, (tag, name, body, featured) in enumerate(packs):
        x = Inches(0.55) + Inches(i * 4.15)
        c = card(s, x, Inches(1.8), Inches(3.95), Inches(4.5), fill=NAVY_CARD)
        if featured:
            top = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, Inches(1.8), Inches(3.95), Inches(0.12))
            fill_shape(top, LIME)
        add_text_box(s, x + Inches(0.3), Inches(2.15), Inches(3.35), Inches(0.35), tag, size=12, bold=True, color=LIME)
        add_text_box(s, x + Inches(0.3), Inches(2.55), Inches(3.35), Inches(0.5), name, size=22, bold=True, color=WHITE)
        add_text_box(s, x + Inches(0.3), Inches(3.25), Inches(3.35), Inches(2.7), body, size=14, color=SOFT)
    footer(s, 12, total)

    # ── 13. Objections ────────────────────────────────────────
    s = new_slide(prs)
    section_title(s, "Objections & réponses", "Anticiper pour closer")
    objs = [
        ("« On a déjà Bloomberg »", "CommoHedge n'est pas un remplacement data desk : c'est le workstation hedge opérationnel (stratégie → book → risk) que Bloomberg seul ne structure pas pour la trésorerie mid-market."),
        ("« Excel nous suffit »", "Excel ne gère pas barriers/MC audités, book multi-users, MTM live, rôles, ni intel freight/géopolitique dans le même flux."),
        ("« Et l'IFRS 9 ? »", "Flags hedge accounting présents — positionner comme support opérationnel, pas moteur comptable complet. Roadmap / intégration ERP possible."),
        ("« Qualité des data ? »", "Terminal futures/options/vol intégré. Pour deals institutionnels : connecteurs licensed en option Enterprise."),
    ]
    for i, (q, a) in enumerate(objs):
        col = i % 2
        row = i // 2
        x = Inches(0.55) + Inches(col * 6.35)
        y = Inches(1.75) + Inches(row * 2.4)
        card(s, x, y, Inches(6.1), Inches(2.2))
        add_text_box(s, x + Inches(0.3), y + Inches(0.25), Inches(5.5), Inches(0.45), q, size=14, bold=True, color=LIME)
        add_text_box(s, x + Inches(0.3), y + Inches(0.85), Inches(5.5), Inches(1.15), a, size=13, color=SOFT)
    footer(s, 13, total)

    # ── 14. Closing CTA ───────────────────────────────────────
    s = new_slide(prs)
    strip = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.18), SLIDE_H)
    fill_shape(strip, LIME)
    add_text_box(s, Inches(0.7), Inches(1.3), Inches(11), Inches(0.4), "CLOSING", size=14, bold=True, color=LIME)
    add_text_box(s, Inches(0.7), Inches(1.85), Inches(11.5), Inches(0.7), "3 phrases à retenir", size=32, bold=True, color=WHITE)

    closings = [
        "Un terminal hedge commodités (et FX) — pas une collection d'outils.",
        "Du pricing exotique au reporting comité, dans le même parcours.",
        "Data marché + freight + intel géopolitique : l'avantage que le mid-market n'a presque jamais.",
    ]
    for i, t in enumerate(closings):
        y = Inches(2.9) + Inches(i * 0.85)
        num = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.75), y, Inches(0.45), Inches(0.45))
        fill_shape(num, LIME)
        add_text_box(s, Inches(0.75), y + Inches(0.05), Inches(0.45), Inches(0.35), str(i + 1), size=14, bold=True, color=LIME_DARK, align=PP_ALIGN.CENTER)
        add_text_box(s, Inches(1.45), y + Inches(0.05), Inches(10.5), Inches(0.5), t, size=16, color=SOFT)

    add_text_box(
        s,
        Inches(0.7),
        Inches(5.7),
        Inches(11),
        Inches(0.5),
        "Prochaine étape : démo guidée 25 min  ·  request access  ·  POC sur votre book",
        size=15,
        bold=True,
        color=LIME,
    )
    footer(s, 14, total)

    prs.save(OUT)
    print(f"Saved: {OUT}")


if __name__ == "__main__":
    build()
